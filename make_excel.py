"""Excel table: expected points of bids 2/3/4 by exact trump holding and off-suit Aces.

Lightly condensed from the every-hand search tables (allhands.py): hands are
grouped by the exact trump cards held and the number of off-suit Aces; all
other side cards are averaged over, weighted by how often each hand is dealt.

    .venv/Scripts/python make_excel.py
"""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ORDER = "23456789TJQKA"
GAMES = (("4-handed", 4, 224), ("6-handed", 6, 112))
TOTAL_HANDS = 20_358_520
FN = "setback_ev_by_trumps_and_aces.xlsx"

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="E4E7EB")
BID_FILLS = {"pass": "F6D5D1", "2": "FDF1C7", "3": "D7EBCE", "4": "9FD08A"}


def table(p, d):
    df = pd.read_csv(f"allhands_{p}p_d{d}_s16.csv.gz")
    hold, ntr, offa = [], [], []
    for h, t in zip(df.hand, df.trump):
        cards = h.split()
        rs = sorted((c[0] for c in cards if c[1] == t), key=ORDER.index, reverse=True)
        hold.append(" ".join("10" if r == "T" else r for r in rs))
        ntr.append(len(rs))
        offa.append(sum(1 for c in cards if c[1] != t and c[0] == "A"))
    df = df.assign(hold=hold, ntr=ntr, offa=offa)
    cols = ["mean", "p2", "p3", "p4", "ev2", "ev3", "ev4"]
    w = df.weight
    agg = df[cols].mul(w, axis=0).assign(weight=w, classes=1, hold=df.hold, ntr=df.ntr, offa=df.offa)
    g = agg.groupby(["hold", "ntr", "offa"]).sum()
    g[cols] = g[cols].div(g.weight, axis=0)
    g["deals"] = g.classes * d
    g = g.reset_index()
    # strongest holdings first: compare trump ranks high to low
    # (a holding sorts before any shorter holding it extends: A K before A)
    g["key"] = [tuple(-ORDER.index("T" if r == "10" else r) for r in h.split()) + (1,) for h in g.hold]
    return g.sort_values(["key", "offa"]).drop(columns=["key", "classes"])


def write_sheet(ws, g, game):
    heads = ["Trump holding", "Trumps", "Off-suit Aces", "Real hands", "Deals simulated",
             "Avg points", "P(make 2)", "P(make 3)", "P(make 4)",
             "EV bid 2", "EV bid 3", "EV bid 4", "Best bid"]
    widths = [17, 8, 12, 12, 14, 11, 11, 11, 11, 10, 10, 10, 10]
    for c, (h, wd) in enumerate(zip(heads, widths), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill = BOLD, HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = wd
    ws.row_dimensions[1].height = 28
    fmts = [None, "0", "0", "#,##0", "#,##0", "0.00", "0.0%", "0.0%", "0.0%",
            "+0.00;-0.00;0.00", "+0.00;-0.00;0.00", "+0.00;-0.00;0.00", None]
    n = len(g)
    for i, r in enumerate(g.itertuples(index=False), start=2):
        vals = [r.hold, r.ntr, r.offa, int(r.weight), int(r.deals), r.mean, r.p2, r.p3, r.p4,
                r.ev2, r.ev3, r.ev4,
                f'=IF(L{i}>=0,"4",IF(K{i}>=0,"3",IF(J{i}>=0,"2","pass")))']
        for c, (v, f) in enumerate(zip(vals, fmts), start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = FONT
            if f:
                cell.number_format = f
            if c >= 2:
                cell.alignment = Alignment(horizontal="center")
    last = n + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:M{last}"
    for col in "JKL":
        ws.conditional_formatting.add(f"{col}2:{col}{last}", ColorScaleRule(
            start_type="num", start_value=-3, start_color="E8A39B",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=3, end_color="7FC06A"))
    for bid, color in BID_FILLS.items():
        ws.conditional_formatting.add(f"M2:M{last}", CellIsRule(
            operator="equal", formula=[f'"{bid}"'], fill=PatternFill("solid", fgColor=color)))
    return n


def write_readme(ws, counts):
    rows = [
        ("Setback: expected points of each bid, by trump holding and off-suit Aces", BOLD),
        ("", None),
        ("What this is", BOLD),
        ("For every combination of trump cards you could hold (and how many Aces you hold in other suits), "
         "how a 2-, 3- or 4-bid works out on average if you win the bid and play the hand.", None),
        ("One sheet per game: 4-handed (two teams of two) and 6-handed (three teams of two).", None),
        ("", None),
        ("Columns", BOLD),
        ("Trump holding: the cards you hold in the suit you name as trump, high to low (10 = ten).", None),
        ("Trumps: how many trumps that is.   Off-suit Aces: Aces you hold in the other suits.", None),
        ("Real hands: how many of the 20,358,520 possible 6-card hands look like this.", None),
        ("Deals simulated: how many deals this row's numbers rest on (more = more precise).", None),
        ("Avg points: your team's average points from the hand (of 4).", None),
        ("P(make k): the chance your team takes at least k points.", None),
        ("EV bid k: expected points at bid k = points you take when you make it, minus k when you're set "
         "(the other team's points are not counted).", None),
        ("Best bid: the highest bid with EV >= 0 (a formula on the EV columns), i.e. the most to go to in the "
         "auction. If a lower bid wins the auction, bid that.", None),
        ("", None),
        ("How the numbers were made", BOLD),
        ("Source: this simulator's every-hand tables (allhands_4p_d224_s16.csv.gz, allhands_6p_d112_s16.csv.gz). "
         "Every distinct 6-card hand was played as bidder against 224 (4-handed) or 112 (6-handed) random deals of "
         "the other cards, by players that look ahead at every decision (16 sampled deals of the hidden cards).", None),
        ("The bidder names trump with fitted weights (trump_weights.py); each row covers the hands that name that "
         "trump. Rows average every hand in the group, weighted by how often it's dealt.", None),
        ("Rules: High, Low (to whoever captures it), Jack, Game (ties score nobody); bids 2-4; the bidder leads "
         "first, which sets trump; you may always trump, and must follow trump when trump is led.", None),
        ("Precision: a single deal's points vary by about 1.5, so a row resting on N deals has an EV error of "
         "about 1.5 / sqrt(N). Common rows are very precise; rows with a few hundred deals are good to about +/-0.1.",
         None),
        ("Not included: 11-bids (sweeps), the auction itself, and the score situation.", None),
        ("", None),
        (f"Rows: 4-handed {counts['4-handed']:,}, 6-handed {counts['6-handed']:,}. "
         "Use the filter arrows in row 1 to find a holding.", None),
    ]
    ws.column_dimensions["A"].width = 130
    for i, (text, font) in enumerate(rows, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font or FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=1, column=1).font = Font(name="Arial", size=14, bold=True)


def main():
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    counts = {}
    for name, p, d in GAMES:
        g = table(p, d)
        assert abs(g.weight.sum() - TOTAL_HANDS) < 1, g.weight.sum()
        counts[name] = write_sheet(wb.create_sheet(name), g, name)
    write_readme(readme, counts)
    wb.calculation.fullCalcOnLoad = True  # Excel computes the Best bid formulas on open
    wb.save(FN)
    print(f"wrote {FN}: " + ", ".join(f"{k} {v:,} rows" for k, v in counts.items()))


if __name__ == "__main__":
    main()
